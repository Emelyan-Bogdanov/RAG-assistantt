import os


def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.txt':
        with open(filepath, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    elif ext == '.pdf':
        try:
            import fitz
            doc = fitz.open(filepath)
            text = ""
            for page in doc:
                text += page.get_text()
            return text
        except ImportError:
            return "PDF extraction requires PyMuPDF. Install: pip install PyMuPDF"
    elif ext == '.docx':
        try:
            from docx import Document
            doc = Document(filepath)
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            return "DOCX extraction requires python-docx. Install: pip install python-docx"
    else:
        return f"[File: {os.path.basename(filepath)}]"


def extract_table(excel: str):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(excel, read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                rows.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(rows)
    except ImportError:
        return "Table extraction requires openpyxl. Install: pip install openpyxl"